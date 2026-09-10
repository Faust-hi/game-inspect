from __future__ import annotations

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocumentClass
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    Image as RLImage,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph as RLParagraph,
    Spacer,
    Table as RLTable,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "docs" / "report" / "Отчет_game-inspect_рабочая_версия.docx"
OUTPUT = ROOT / "docs" / "report" / "Отчет_game-inspect_рабочая_версия.pdf"

FONT_DIR = Path(r"C:\Windows\Fonts")
pdfmetrics.registerFont(TTFont("ReportTimes", str(FONT_DIR / "times.ttf")))
pdfmetrics.registerFont(TTFont("ReportTimes-Bold", str(FONT_DIR / "timesbd.ttf")))
pdfmetrics.registerFont(TTFont("ReportTimes-Italic", str(FONT_DIR / "timesi.ttf")))
pdfmetrics.registerFont(TTFont("ReportTimes-BoldItalic", str(FONT_DIR / "timesbi.ttf")))
pdfmetrics.registerFontFamily("ReportTimes", normal="ReportTimes", bold="ReportTimes-Bold", italic="ReportTimes-Italic", boldItalic="ReportTimes-BoldItalic")


PAGE_W, PAGE_H = A4
LEFT = 30 * mm
RIGHT = 10 * mm
TOP = 20 * mm
BOTTOM = 20 * mm
USABLE_W = PAGE_W - LEFT - RIGHT


def iter_blocks(parent):
    if isinstance(parent, DocumentClass):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_text(p: Paragraph) -> str:
    return "".join(node.text or "" for node in p._p.iter(qn("w:t")))


def has_page_break(p: Paragraph) -> bool:
    for br in p._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def has_page_break_before(p: Paragraph) -> bool:
    """Return the DOCX pageBreakBefore flag used for major sections."""
    try:
        return bool(p.paragraph_format.page_break_before)
    except Exception:
        return False


def paragraph_images(p: Paragraph):
    """Yield embedded images from a DOCX paragraph with their native aspect ratio."""
    for blip in p._p.iter(qn("a:blip")):
        rel_id = blip.get(qn("r:embed"))
        if not rel_id:
            continue
        try:
            part = p.part.related_parts[rel_id]
            blob = part.blob
        except Exception:
            continue
        width = None
        height = None
        inline = blip.getparent()
        while inline is not None:
            extent = inline.find(qn("wp:extent"))
            if extent is not None:
                try:
                    width = int(extent.get("cx")) / 914400 * 72
                    height = int(extent.get("cy")) / 914400 * 72
                except (TypeError, ValueError):
                    pass
                break
            inline = inline.getparent()
        if not width or not height:
            width, height = USABLE_W, USABLE_W * 0.4
        scale = min(1.0, USABLE_W / width)
        yield BytesIO(blob), width * scale, height * scale


def html_text(text: str) -> str:
    return escape(text).replace("\n", "<br/>")


class ReportDocTemplate(BaseDocTemplate):
    def __init__(self, filename: str, **kwargs):
        super().__init__(filename, pagesize=A4, leftMargin=LEFT, rightMargin=RIGHT, topMargin=TOP, bottomMargin=BOTTOM, **kwargs)
        frame = Frame(LEFT, BOTTOM, USABLE_W, PAGE_H - TOP - BOTTOM, id="normal")
        self.addPageTemplates([PageTemplate(id="report", frames=[frame], onPage=draw_page_number)])
        self._heading_index = 0

    def afterFlowable(self, flowable: Flowable) -> None:
        if not isinstance(flowable, RLParagraph):
            return
        style_name = flowable.style.name
        if not style_name.startswith("ReportHeading"):
            return
        level = int(style_name[-1])
        text = flowable.getPlainText()
        self._heading_index += 1
        key = f"heading-{self._heading_index}"
        self.canv.bookmarkPage(key)
        self.canv.addOutlineEntry(text, key, level=level - 1, closed=False)
        self.notify("TOCEntry", (level, text, self.page))


def draw_page_number(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("ReportTimes", 12)
    canvas.drawCentredString(PAGE_W / 2, 10 * mm, str(doc.page))
    canvas.restoreState()


def styles():
    return {
        "body": ParagraphStyle(
            "ReportBody", fontName="ReportTimes", fontSize=14, leading=21,
            alignment=TA_JUSTIFY, firstLineIndent=12.5 * mm, spaceBefore=0, spaceAfter=0,
        ),
        "body_noindent": ParagraphStyle(
            "ReportBodyNoIndent", fontName="ReportTimes", fontSize=14, leading=21,
            alignment=TA_JUSTIFY, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "bullet": ParagraphStyle(
            "ReportBullet", fontName="ReportTimes", fontSize=14, leading=21,
            leftIndent=12.5 * mm, firstLineIndent=0, bulletIndent=5 * mm, alignment=TA_JUSTIFY,
            spaceBefore=0, spaceAfter=0,
        ),
        "h1": ParagraphStyle(
            "ReportHeading1", fontName="ReportTimes-Bold", fontSize=18, leading=23,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=10, spaceAfter=6, keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "ReportHeading2", fontName="ReportTimes-Bold", fontSize=16, leading=20,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=8, spaceAfter=5, keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "ReportHeading3", fontName="ReportTimes-Bold", fontSize=14, leading=18,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=6, spaceAfter=4, keepWithNext=True,
        ),
        "caption": ParagraphStyle(
            "ReportCaption", fontName="ReportTimes-Italic", fontSize=12, leading=15,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=5, spaceAfter=4, keepWithNext=True,
        ),
        "equation": ParagraphStyle(
            "ReportEquation", fontName="ReportTimes", fontSize=13, leading=18,
            alignment=TA_CENTER, firstLineIndent=0, spaceBefore=4, spaceAfter=4,
        ),
        "title": ParagraphStyle(
            "ReportTitle", fontName="ReportTimes-Bold", fontSize=16, leading=21,
            alignment=TA_CENTER, firstLineIndent=0, spaceBefore=0, spaceAfter=4,
        ),
        "cover": ParagraphStyle(
            "ReportCover", fontName="ReportTimes-Bold", fontSize=14, leading=18,
            alignment=TA_CENTER, firstLineIndent=0, spaceBefore=0, spaceAfter=4,
        ),
        "cover_right": ParagraphStyle(
            "ReportCoverRight", fontName="ReportTimes", fontSize=14, leading=21,
            alignment=TA_RIGHT, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "toc": ParagraphStyle(
            "ReportTOC", fontName="ReportTimes", fontSize=12, leading=16,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "toc0": ParagraphStyle(
            "ReportTOC0", fontName="ReportTimes-Bold", fontSize=12, leading=16,
            alignment=TA_LEFT, firstLineIndent=0, leftIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "toc1": ParagraphStyle(
            "ReportTOC1", fontName="ReportTimes", fontSize=12, leading=16,
            alignment=TA_LEFT, firstLineIndent=0, leftIndent=8 * mm, spaceBefore=0, spaceAfter=0,
        ),
        "toc2": ParagraphStyle(
            "ReportTOC2", fontName="ReportTimes", fontSize=11, leading=15,
            alignment=TA_LEFT, firstLineIndent=0, leftIndent=16 * mm, spaceBefore=0, spaceAfter=0,
        ),
        "source": ParagraphStyle(
            "ReportSource", fontName="ReportTimes", fontSize=11, leading=14,
            alignment=TA_LEFT, leftIndent=12.5 * mm, firstLineIndent=-12.5 * mm, spaceAfter=4,
        ),
        "code": ParagraphStyle(
            "ReportCode", fontName="Courier", fontSize=9, leading=11,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "table": ParagraphStyle(
            "ReportTable", fontName="ReportTimes", fontSize=9.5, leading=11.5,
            alignment=TA_LEFT, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
        "table_header": ParagraphStyle(
            "ReportTableHeader", fontName="ReportTimes-Bold", fontSize=9.5, leading=11.5,
            alignment=TA_CENTER, firstLineIndent=0, spaceBefore=0, spaceAfter=0,
        ),
    }


def table_widths(table: Table) -> list[float]:
    raw = []
    for cell in table.rows[0].cells:
        try:
            raw.append(float(cell.width) / 360000 * mm)
        except Exception:
            raw.append(1.0)
    if not raw or sum(raw) <= 0:
        return [USABLE_W / len(table.columns)] * len(table.columns)
    scale = USABLE_W / sum(raw)
    return [x * scale for x in raw]


def make_table(table: Table, st) -> RLTable:
    if len(table.columns) == 1 and len(table.rows) == 1:
        value = "\n".join(p.text for p in table.cell(0, 0).paragraphs).strip()
        code = RLParagraph(html_text(value).replace("\n", "<br/>"), st["code"])
        result = RLTable([[code]], colWidths=[USABLE_W], hAlign="CENTER")
        result.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#404040")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F5F7")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return result
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            value = "\n".join(p.text for p in cell.paragraphs).strip()
            style = st["table_header"] if row_index == 0 else st["table"]
            cells.append(RLParagraph(html_text(value), style))
        rows.append(cells)
    widths = table_widths(table)
    result = RLTable(rows, colWidths=widths, repeatRows=1, hAlign="CENTER")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#404040")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for i in range(1, len(rows)):
        if i % 2 == 0:
            commands.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F3F5F7")))
    result.setStyle(TableStyle(commands))
    return result


def is_top_heading(text: str) -> bool:
    return (
        text[:2].isdigit()
        or text in {"ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"}
        or text.startswith("ПРИЛОЖЕНИЕ ")
    )


def build() -> None:
    doc = Document(str(INPUT))
    st = styles()
    story = []
    title_page = True
    toc_added = False
    pending_toc = False

    for block in iter_blocks(doc):
        if isinstance(block, Table):
            story.append(make_table(block, st))
            story.append(Spacer(1, 2))
            continue

        if has_page_break_before(block) and not title_page:
            story.append(PageBreak())
        images = list(paragraph_images(block))
        if images:
            for blob, width, height in images:
                story.append(RLImage(blob, width=width, height=height))
                story.append(Spacer(1, 2))
            continue

        text = paragraph_text(block)
        if has_page_break(block):
            story.append(PageBreak())
            title_page = False
            continue
        stripped = text.strip()
        if not stripped:
            story.append(Spacer(1, 5 if not title_page else 10))
            continue
        style_name = block.style.name if block.style else "Normal"

        if pending_toc and "Содержание будет обновлено" in stripped:
            story.append(TableOfContents(levelStyles=[st["toc0"], st["toc1"], st["toc2"]], dotsMinLevel=0))
            toc_added = True
            pending_toc = False
            continue
        if "Содержание будет обновлено" in stripped or stripped.startswith("[При необходимости обновить поля содержания"):
            continue

        if style_name.startswith("Heading 1"):
            if stripped == "СОДЕРЖАНИЕ":
                story.append(RLParagraph(stripped, st["h1"]))
                pending_toc = True
            else:
                story.append(RLParagraph(html_text(stripped), st["h1"]))
            continue
        if style_name.startswith("Heading 2"):
            story.append(RLParagraph(html_text(stripped), st["h2"]))
            continue
        if style_name.startswith("Heading 3"):
            story.append(RLParagraph(html_text(stripped), st["h3"]))
            continue
        if stripped.startswith("Таблица ") or (
            stripped.startswith("Листинг ") and len(stripped) > 8 and stripped[8].isdigit()
        ):
            story.append(RLParagraph(html_text(stripped), st["caption"]))
            continue
        if stripped.startswith("[Поля ") or stripped.startswith("[После очистки"):
            story.append(RLParagraph(f'<font color="#6B7280"><i>{html_text(stripped)}</i></font>', st["body_noindent"]))
            continue

        # Cover text is intentionally centered/right-aligned to match the DOCX.
        if title_page:
            if stripped.startswith(("Выполнил:", "Группа:", "Руководитель:")):
                story.append(RLParagraph(html_text(stripped), st["cover_right"]))
            elif "ОТЧЁТ ПО КУРСОВОМУ ПРОЕКТУ" in stripped or stripped.startswith("ИНФОРМАЦИОННАЯ СИСТЕМА"):
                story.append(RLParagraph(html_text(stripped), st["title"]))
            else:
                story.append(RLParagraph(html_text(stripped), st["cover"]))
            continue

        if stripped.startswith("Ключевые слова:"):
            story.append(RLParagraph(html_text(stripped), st["body_noindent"]))
        elif stripped.startswith("[Здесь можно") or stripped.startswith("[Поля в квадратных"):
            story.append(RLParagraph(f'<font color="#6B7280"><i>{html_text(stripped)}</i></font>', st["body_noindent"]))
        elif stripped.startswith("rᵢ") or stripped.startswith("Cᵢ") or stripped.startswith("Bкадр") or stripped.startswith("Lкадр"):
            story.append(RLParagraph(html_text(stripped), st["equation"]))
        elif style_name == "List Bullet":
            story.append(RLParagraph(html_text(stripped), st["bullet"], bulletText="•"))
        elif stripped.startswith("[") and "] " in stripped and "http" in stripped:
            story.append(RLParagraph(html_text(stripped), st["source"]))
        else:
            story.append(RLParagraph(html_text(stripped), st["body"]))

    if not toc_added and pending_toc:
        story.append(RLParagraph("Содержание будет обновлено при открытии документа в Word", st["toc"]))

    class _Keep(Flowable):
        pass

    pdf = ReportDocTemplate(str(OUTPUT), title="Отчёт по проекту game-inspect", author="[ФИО студента]")
    pdf.multiBuild(story)
    print(OUTPUT)


if __name__ == "__main__":
    build()
