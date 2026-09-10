from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "report"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_DOCX = OUT_DIR / "Отчет_game-inspect_рабочая_версия.docx"

BLACK = "000000"
TABLE_BORDER = "D9D9D9"
TABLE_HEADER = "404040"
TABLE_ALT = "F3F5F7"
PLACEHOLDER = "6B7280"


def set_run_font(run, name: str = "Times New Roman", size: float = 14, *, bold: bool | None = None,
                 italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def set_style_font(style, name: str, size: float, *, bold: bool = False, italic: bool = False, color: str = BLACK) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = TABLE_BORDER, size: str = "6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_mm: float) -> None:
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_mm * 56.7)))
    tc_w.set(qn("w:type"), "dxa")


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Mm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run()
    set_run_font(run, size=12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Содержание будет обновлено при открытии документа в Word"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_placeholder(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, size=12, italic=True, color=PLACEHOLDER)


def add_table(doc: Document, caption: str, headers: Sequence[str], widths: Sequence[float],
              rows: Sequence[Sequence[str]], *, font_size: float = 9.5) -> None:
    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Mm(0)
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    run = cap.add_run(caption)
    set_run_font(run, size=12, italic=True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_margins(cell)
        set_cell_shading(cell, TABLE_HEADER)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, TABLE_ALT)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)

    after = doc.add_paragraph()
    after.paragraph_format.first_line_indent = Mm(0)
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.line_spacing = 1.0


def add_paragraph(doc: Document, text: str = "", *, first_line: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  size: float = 14, bold: bool = False, italic: bool = False, keep_next: bool = False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Mm(12.5) if first_line else Mm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = keep_next
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Mm(12.5)
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(item)
        set_run_font(r, size=14)


def add_equation(doc: Document, text: str, number: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{text}    {number}")
    set_run_font(r, size=13)


_TEMP_ASSETS: list[Path] = []


def add_code_listing(doc: Document, caption: str, code: str) -> None:
    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Mm(0)
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    set_run_font(r, size=12, italic=True)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=BLACK, size="8")
    cell = table.cell(0, 0)
    set_cell_width(cell, 155)
    set_cell_margins(cell, top=90, start=130, bottom=90, end=130)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    for line_index, line in enumerate(code.strip("\n").splitlines()):
        run = p.add_run(line)
        set_run_font(run, name="Courier New", size=9)
        if line_index < len(code.strip("\n").splitlines()) - 1:
            run.add_break()

    after = doc.add_paragraph()
    after.paragraph_format.first_line_indent = Mm(0)
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.line_spacing = 1.0


def add_pipeline_figure(doc: Document) -> None:
    """Add one compact architecture diagram without introducing external graphics."""
    from PIL import Image, ImageDraw, ImageFont

    normal_path = r"C:\Windows\Fonts\times.ttf"
    bold_path = r"C:\Windows\Fonts\timesbd.ttf"
    normal = ImageFont.truetype(normal_path, 30)
    small = ImageFont.truetype(normal_path, 25)
    bold = ImageFont.truetype(bold_path, 29)
    image = Image.new("RGB", (1600, 640), "white")
    draw = ImageDraw.Draw(image)

    def box(x0, y0, x1, y1, lines, *, fill="#F3F5F7", font=normal):
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=fill, outline="#404040", width=3)
        text = "\n".join(lines)
        draw.multiline_text(((x0 + x1) / 2, (y0 + y1) / 2), text, font=font, fill="#000000",
                            anchor="mm", align="center", spacing=7)

    def arrow(x0, y0, x1, y1):
        draw.line((x0, y0, x1, y1), fill="#404040", width=4)
        if x1 >= x0:
            draw.polygon([(x1, y1), (x1 - 18, y1 - 10), (x1 - 18, y1 + 10)], fill="#404040")
        else:
            draw.polygon([(x1, y1), (x1 + 18, y1 - 10), (x1 + 18, y1 + 10)], fill="#404040")

    top = [
        (35, 115, 285, 245, ["Профиль", "проекта"]),
        (345, 115, 595, 245, ["FastAPI", "валидация"]),
        (655, 115, 905, 245, ["Доменное", "ядро расчёта"]),
        (965, 115, 1215, 245, ["Рекомендации", "и корзина"]),
        (1275, 115, 1565, 245, ["React", "интерфейс"]),
    ]
    for args in top:
        box(*args)
    for left, right in zip(top, top[1:]):
        arrow(left[2] + 10, 180, right[0] - 12, 180)

    box(400, 380, 700, 505, ["SQLite", "каталог"], fill="#FFFFFF", font=small)
    box(900, 380, 1200, 505, ["sessionStorage", "текущий проект"], fill="#FFFFFF", font=small)
    arrow(550, 380, 780, 250)
    arrow(1000, 380, 1415, 250)
    draw.text((800, 40), "Упрощённый конвейер расчёта", font=bold, fill="#000000", anchor="ma")

    with NamedTemporaryFile(prefix="game_inspect_pipeline_", suffix=".png", delete=False) as handle:
        path = Path(handle.name)
    image.save(path)
    _TEMP_ASSETS.append(path)
    doc.add_picture(str(path), width=Mm(155))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Mm(0)
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run("Рисунок 3.1 - Упрощённый конвейер формирования результата")
    set_run_font(r, size=12, bold=True)


def add_heading(doc: Document, text: str, level: int = 1, *, page_break: bool = False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    return p


def add_source(doc: Document, number: int, title: str, url: str, note: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(12.5)
    p.paragraph_format.first_line_indent = Mm(-12.5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"[{number}] {title}. ")
    set_run_font(r, size=11)
    add_hyperlink(p, url, url)
    r2 = p.add_run(f" {note}")
    set_run_font(r2, size=11)


doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.left_margin = Mm(30)
section.right_margin = Mm(10)
section.top_margin = Mm(20)
section.bottom_margin = Mm(20)
section.header_distance = Mm(10)
section.footer_distance = Mm(10)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, "Times New Roman", 14)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.first_line_indent = Mm(12.5)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)

for name, size, bold in (("Heading 1", 18, True), ("Heading 2", 16, True), ("Heading 3", 14, True)):
    style = styles[name]
    set_style_font(style, "Times New Roman", size, bold=bold)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.first_line_indent = Mm(0)
    style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

if "Caption Report" not in [s.name for s in styles]:
    caption_style = styles.add_style("Caption Report", WD_STYLE_TYPE.PARAGRAPH)
else:
    caption_style = styles["Caption Report"]
set_style_font(caption_style, "Times New Roman", 12, italic=True)

for style_name in ("List Bullet", "List Number"):
    if style_name in styles:
        style = styles[style_name]
        set_style_font(style, "Times New Roman", 14)
        style.paragraph_format.line_spacing = 1.5

doc.core_properties.title = "Информационная система поддержки принятия решений по оптимизации разработки компьютерных игр"
doc.core_properties.subject = "Отчёт по проекту game-inspect"
doc.core_properties.author = "[ФИО студента]"
doc.core_properties.comments = "Рабочая версия для финальной редакции"

# Обновление полей Word при открытии документа.
settings = doc.settings.element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

# Нумерация страниц.
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Mm(0)
fp.paragraph_format.space_before = Pt(0)
fp.paragraph_format.space_after = Pt(0)
add_page_field(fp)

# Титульный лист.
for _ in range(3):
    add_paragraph(doc, "", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
for text in ("[НАИМЕНОВАНИЕ ОБРАЗОВАТЕЛЬНОЙ ОРГАНИЗАЦИИ]", "[ФАКУЛЬТЕТ / ИНСТИТУТ]", "[КАФЕДРА]"):
    p = add_paragraph(doc, text, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    p.paragraph_format.space_after = Pt(4)
add_paragraph(doc, "", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph(doc, "ОТЧЁТ ПО КУРСОВОМУ ПРОЕКТУ", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True)
add_paragraph(doc, "", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
title = add_paragraph(
    doc,
    "ИНФОРМАЦИОННАЯ СИСТЕМА ПОДДЕРЖКИ ПРИНЯТИЯ РЕШЕНИЙ ПО ОПТИМИЗАЦИИ РАЗРАБОТКИ КОМПЬЮТЕРНЫХ ИГР",
    first_line=False,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    size=16,
    bold=True,
)
title.paragraph_format.keep_together = True
for _ in range(5):
    add_paragraph(doc, "", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
for text in (
    "Выполнил: [ФИО студента]",
    "Группа: [номер группы]",
    "Руководитель: [ФИО, должность]",
):
    add_paragraph(doc, text, first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT, size=14)
for _ in range(4):
    add_paragraph(doc, "", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph(doc, "[ГОРОД]", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
add_paragraph(doc, "2026", first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
doc.add_page_break()

# Аннотация и содержание.
add_heading(doc, "АННОТАЦИЯ", 1)
add_paragraph(doc, "Отчёт посвящён разработке локальной информационной системы поддержки принятия решений для выбора технических способов реализации функций компьютерной игры. Система принимает профиль проекта, подбирает применимые методы, ранжирует их по нескольким критериям, показывает связь с инструментами игровых движков, выявляет ограничения и рассчитывает ориентировочный класс аппаратного обеспечения.")
add_paragraph(doc, "В работе зафиксирована граница предметной области: рассматриваются решения, влияющие на архитектуру, вычислительную нагрузку или потребление ресурсов. Монетизация, дата выхода, сюжетные развилки и количество концовок в расчёт не включаются. Сохраняются геймплейные решения, которые создают технические требования, в том числе разрушаемость окружения, масштаб мира, количество NPC и мультиплеер.")
add_paragraph(doc, "Основная функциональность проекта реализована. Текущая версия содержит каталог технических методов, многокритериальное ранжирование TOPSIS, расчёт ресурсов, корзину выбранных решений, объяснение результатов, административное наполнение и набор автоматических проверок. В документе оставлены поля для обновления после финальной очистки тестов, окончательной проверки быстрого старта и уточнения отдельных спорных записей каталога.")
p = add_paragraph(doc, "Ключевые слова: ", first_line=False, size=12, bold=True)
r = p.add_run("поддержка принятия решений, разработка компьютерных игр, техническая оптимизация, TOPSIS, аппаратная оценка, игровые движки, производительность.")
set_run_font(r, size=12)
doc.add_page_break()

add_heading(doc, "СОДЕРЖАНИЕ", 1)
toc = doc.add_paragraph()
add_toc_field(toc)
add_paragraph(doc, "", first_line=False)
add_placeholder(add_paragraph(doc, "", first_line=False), "[При необходимости обновить поля содержания в Word сочетанием Ctrl+A, F9.]")
doc.add_page_break()

add_heading(doc, "СОКРАЩЕНИЯ И ОБОЗНАЧЕНИЯ", 1)
add_table(doc, "Таблица 0.1 - Основные сокращения", ["Сокращение", "Расшифровка", "Использование в работе"], [28, 58, 84], [
    ("API", "Application Programming Interface", "программный интерфейс backend"),
    ("CPU", "Central Processing Unit", "центральный процессор"),
    ("GPU", "Graphics Processing Unit", "графический процессор"),
    ("RAM", "оперативная память", "память системы"),
    ("VRAM", "видеопамять", "память графического устройства"),
    ("RT", "Ray Tracing", "трассировка лучей"),
    ("LOD / HLOD", "Level of Detail / Hierarchical LOD", "уровни детализации"),
    ("NPC", "Non-Player Character", "неигровой персонаж"),
    ("TOPSIS", "Technique for Order Preference by Similarity to Ideal Solution", "многокритериальное ранжирование"),
]
)
doc.add_page_break()

# Введение.
add_heading(doc, "ВВЕДЕНИЕ", 1)
add_paragraph(doc, "Производительность компьютерной игры определяется не одной настройкой, а взаимодействием нескольких подсистем: подготовки команд рендера, растеризации, освещения, физики, искусственного интеллекта, анимации, потоковой загрузки, памяти и сетевой синхронизации. Поэтому решение, принятое на стадии проектирования, может повлиять на структуру данных, набор инструментов движка, порядок подготовки контента и требования к оборудованию. Задача системы состоит в том, чтобы сделать такие зависимости видимыми до появления готового игрового билда.")
add_paragraph(doc, "Актуальность проекта определяется тем, что существующие средства обычно решают отдельные задачи. Документация движка описывает конкретный механизм, профилировщик измеряет уже созданный фрагмент игры, а каталог оборудования сообщает характеристики устройств. Между этими источниками остаётся практическая задача выбора: какой способ реализации соответствует функциям и ограничениям конкретного проекта, какие компромиссы он создаёт и что произойдёт при замене решения в уже сформированном наборе.")
add_paragraph(doc, "Цель проекта - разработать информационную систему поддержки принятия решений, которая связывает профиль игры, технические функции, варианты их реализации, совместимость решений и ориентировочные требования к ресурсам. Система должна выдавать объяснимый результат и явно отделять расчётную модель от измеренных данных.")
add_paragraph(doc, "Объект исследования - технические решения, принимаемые при проектировании и разработке компьютерных игр. Предмет исследования - модели представления таких решений, правила их применимости и алгоритмы ранжирования с учётом влияния на архитектуру, вычислительную нагрузку и ресурсы.")
add_paragraph(doc, "Для достижения цели решены следующие задачи:")
add_bullets(doc, [
    "определена предметная область и выделены входные параметры профиля проекта;",
    "сформирован каталог игровых функций, технических методов, движков, инструментов и аппаратных записей;",
    "реализованы фильтрация недопустимых решений и многокритериальное ранжирование применимых вариантов;",
    "предусмотрены объяснения результата, условия применения, конфликты, зависимости и взаимодополняющие связи;",
    "разработана ориентировочная модель стоимости кадра, памяти и подбора референсного класса оборудования;",
    "реализована корзина текущих решений с пересчётом результата после изменения профиля или состава набора;",
    "проведены проверки чувствительности к входным параметрам, ограничений, сохранения и защиты от устаревших результатов.",
])
add_paragraph(doc, "Методы работы включают декомпозицию игровой нагрузки по подсистемам, формализацию правил предметной области, многокритериальный метод TOPSIS, анализ связей между решениями, тестирование граничных и сценарных случаев, а также проверку воспроизводимости расчёта. Коэффициенты нагрузки в текущей реализации являются экспертными оценками; они не заменяют профилирование конкретной игры.")
add_paragraph(doc, "Практическая значимость заключается в том, что система позволяет сформировать обоснованный предварительный план технических решений до интеграции их в игровой проект. Результат используется как средство анализа и подготовки прототипа, а не как автоматическое доказательство достижения заданного FPS.")

# Section 1.
add_heading(doc, "1 АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ И ПОСТАНОВКА ЗАДАЧИ", 1, page_break=True)
add_heading(doc, "1.1 Особенности выбора технических решений", 2)
add_paragraph(doc, "В разработке игры одна и та же функция может быть реализована несколькими способами. Например, потоковая загрузка может опираться на разбиение мира, кэширование страниц или виртуальное текстурирование; управление детализацией может использовать отдельные уровни мешей, иерархические группы или импосторы. Способ реализации меняет не только локальную стоимость кадра, но и требования к подготовке контента, памяти, загрузке, инструментам движка и проверке качества.")
add_paragraph(doc, "Важным является различие между функцией игры и методом её реализации. Функция отвечает на вопрос, что должно работать: открытый мир, физика разрушения, толпа NPC или сетевое взаимодействие. Метод отвечает на вопрос, как это организовать: culling, LOD, batching, стриминг, ограничение частоты обновления, репликация или другой механизм. Такая пара позволяет связать пользовательское требование с инженерным решением.")
add_paragraph(doc, "Для анализа стадий разработки использованы идеи разделения проектных параметров, измеряемых характеристик и регрессионных зависимостей, представленные в исследовании GAMORRA [1]. Коэффициенты из внешней работы в систему не переносятся: различаются входные параметры, наборы игр и аппаратные условия. Источник используется как основание для разделения модели и измерительного протокола.")

add_heading(doc, "1.2 Пользовательские роли и потребности системы", 2)
add_paragraph(doc, "Пользователь системы - команда или разработчик, которому нужно принять техническое решение на стадии концепции, прототипа или производства. Роли ниже являются функциональными: в небольшой команде их может совмещать один человек.")
add_table(doc, "Таблица 1.1 - Роли пользователей системы", ["Роль", "Что требуется решить", "Результат работы"], [43, 67, 60], [
    ("Технический руководитель", "выбрать архитектурный вариант и ограничения", "приоритетный набор методов и аппаратный ориентир"),
    ("Программист или technical artist", "сопоставить метод с движком, контентом и пайплайном", "условия применения, риски качества и интеграции"),
    ("Аналитик проекта", "проверить профиль, корзину и объяснимость результата", "воспроизводимый расчёт и список вопросов для проверки"),
])
add_paragraph(doc, "Ключевые требования к системе сформулированы в проверяемом виде:")
add_bullets(doc, [
    "FR-01: принять профиль проекта, нормализовать его и явно обработать неизвестные значения;",
    "FR-02: отфильтровать неприменимые методы с объяснением причины и ранжировать допустимые варианты;",
    "FR-03: рассчитать текущую корзину, конфликты, зависимости и сводное влияние ресурсов;",
    "FR-04: подобрать референсный класс оборудования с учётом производительности, памяти и обязательных возможностей;",
    "NFR-01: для одного профиля и одной корзины выдавать воспроизводимый отпечаток входа и не показывать результат старого состояния.",
])
add_paragraph(doc, "Критерии качества дополняют требования: результат должен быть объяснимым, изменения значимых входов должны отражаться в расчёте, а экспертные оценки и модельные примеры должны быть отделены от измеренных данных.")
add_paragraph(doc, "Таким образом, потребности пользователя покрываются анкетой профиля, фильтрацией и ранжированием каталога, карточкой последствий, корзиной, аппаратной оценкой и сбросом устаревшего результата. Эти функции образуют единый проверяемый сценарий.")

add_heading(doc, "1.3 Анализ существующих средств и границы их применимости", 2)
add_paragraph(doc, "Документация игровых движков является необходимым источником технических сведений, но обычно описывает отдельный инструмент и его ограничения. Она не выполняет сопоставление методов разных движков по критериям проекта, не формирует текущую корзину и не рассчитывает общий профиль ресурсов.")
add_paragraph(doc, "Профилировщики и средства захвата времени кадра полезны после появления работающей сцены. Они позволяют измерять CPU-, GPU- и display-времена, но требуют готового билда или инструментированного прототипа. В рамках проекта такие средства рассматриваются как внешняя проверка выбранного решения, а не как встроенная часть DSS. Это принципиально ограничивает область обещаний системы.")
add_paragraph(doc, "Каталоги оборудования и минимальные требования игр дают ориентир по устройствам, но не объясняют, какие подсистемы создают нагрузку и как меняется результат при замене метода. Кроме того, официальное требование часто относится к конкретному пресету, разрешению и версии игры. Поэтому в проекте аппаратный каталог используется для подбора референсного класса, а не для выдачи гарантии производительности.")
add_paragraph(doc, "Обобщённые системы архитектурного моделирования описывают потоки работ и ресурсы, но для текущей задачи нужна компактная модель с игровыми функциями, методами и условиями совместимости. Универсальное хранилище всех доказательств в границы проекта не входит.")

add_heading(doc, "1.4 Границы предметной области", 2)
add_table(doc, "Таблица 1.2 - Правило включения решений", ["Категория", "Решение", "Обоснование"], [36, 58, 76], [
    ("Включается", "разрушаемость окружения", "влияет на физику, динамическую геометрию, кэширование и сеть"),
    ("Включается", "масштаб мира", "влияет на стриминг, память, подготовку контента и резидентность"),
    ("Включается", "количество NPC", "влияет на AI, анимацию, симуляцию и сетевую репликацию"),
    ("Включается", "мультиплеер", "задаёт сетевую модель, частоты, роли сервера и требования к детерминизму"),
    ("Включается", "разрешение, FPS, качество", "меняют бюджет кадра и стоимость соответствующих проходов"),
    ("Исключается", "монетизация", "не является параметром вычислительной архитектуры в данной задаче"),
    ("Исключается", "дата выхода", "не задаёт физическую стоимость выполнения игры"),
    ("Исключается", "сюжетные развилки и количество концовок", "не создают сами по себе проверяемую ресурсную нагрузку"),
])
add_paragraph(doc, "Критерий включения сформулирован как влияние на архитектуру, вычислительную нагрузку или потребление ресурсов. Это не утверждение о незначимости исключённых решений для игры в целом; они просто находятся за пределами рассматриваемой технической модели.")

add_heading(doc, "1.5 Постановка задачи и критерии качества", 2)
add_paragraph(doc, "Система должна принимать профиль игры и возвращать набор применимых технических методов, ранжированный по выбранному приоритету. Для каждой записи требуется отображать причины выбора, условия применения, предполагаемые изменения ресурсов, влияние на качество и стадию внедрения. Неприменимые записи не должны исчезать бесследно: причина исключения должна быть видна пользователю.")
add_paragraph(doc, "Критерии качества: детерминированность, объяснимость, согласованность экранов, отсутствие неподтверждённых обещаний и реакция на значимый вход. После изменения профиля или корзины старый результат не должен оставаться видимым для нового состояния.")
add_paragraph(doc, "Основная функциональность по этим критериям реализована. Финальная чистка тестов, быстрый старт и дополнительная проверка каталога вынесены в приложение В.")

# Section 2.
add_heading(doc, "2 МОДЕЛЬ И МЕТОДЫ ПОДДЕРЖКИ ПРИНЯТИЯ РЕШЕНИЙ", 1, page_break=True)
add_heading(doc, "2.1 Профиль проекта и технические входы", 2)
add_paragraph(doc, "Профиль ProjectProfile содержит 41 поле. Поля сгруппированы по смыслу, чтобы не смешивать идентификацию проекта, требования геймплея, параметры активной сцены, ограничения ресурсов и политику выбора. Наличие поля в анкете не означает, что оно является измеренным параметром: для неизвестного значения система должна показывать допущение или снижать уверенность.")
add_table(doc, "Таблица 2.1 - Группы входных параметров профиля", ["Группа", "Примеры полей", "Роль в расчёте"], [39, 70, 61], [
    ("Контекст", "format, world_type, scale, stage, engine, engine_version, platforms", "применимость, платформенные и стадийные условия"),
    ("Активная сцена", "object_count_level, object_count, npc_count_level, npc_count, local_view_count", "работа на кадр и оценка уверенности"),
    ("Геймплейные требования", "functions, multiplayer, player_count, simulation_radius_m", "выбор подсистем и архитектурных ограничений"),
    ("Рендеринг", "target_resolution, target_quality, target_fps, render_api", "бюджет кадра и стоимость пиксельных стадий"),
    ("Симуляция", "physics_tick_hz, audio_complexity, network_topology", "отдельные частоты и виды нагрузки"),
    ("Память и данные", "memory_model, streaming_pool_gb, storage_type, size_limit_gb", "состав памяти, потоковая загрузка и ограничения"),
    ("Политика выбора", "complexity_tolerance, deadline_weeks, priority, cpu_budget", "фильтры, веса и риск проекта"),
])
add_paragraph(doc, "Некоторые поля требуют аккуратного определения. Количество объектов не должно автоматически трактоваться как количество видимых объектов, а количество NPC - как число одновременно симулируемых агентов. Число игроков не равно числу камер и не задаёт топологию сети. Стадия разработки влияет на своевременность внедрения и стоимость поздней переработки, но не должна менять физическую стоимость одного и того же кадра без изменения реализации.")

add_heading(doc, "2.2 Паспорт технического метода", 2)
add_paragraph(doc, "Запись каталога описывает не только название оптимизации. В ней фиксируются функция или общий уровень применения, вид решения, область эффекта, ограничения формата и движка, ресурсные эффекты, влияние на качество и концепцию, стоимость внедрения, стоимость позднего внедрения, рекомендуемая стадия, шаги реализации, примеры использования и источник.")
add_paragraph(doc, "Основная причинная цепочка в паспорте имеет вид: игровая функция -> базовый способ реализации -> заменяемая или сокращаемая работа -> технический метод -> затронутые ресурсы и проверяемые последствия. Если промежуточный механизм не определён, утверждение о причинном эффекте заменяется на краткое допущение или оставляется пустым.")
add_paragraph(doc, "В паспорте различаются четыре статуса: механизм, подтверждённый источником; экспертное допущение для численного коэффициента; измеренный результат конкретного стенда; открытый вопрос, для которого поле оставляется пустым или сопровождается короткой оговоркой.")

add_heading(doc, "2.3 Фильтрация и оценка применимости", 2)
add_paragraph(doc, "Сначала система формирует кандидатов по выбранным функциям и добавляет общие методы. Затем для каждого метода выполняется проверка применимости. Жёсткими основаниями исключения являются несовместимый формат, тип мира, платформа, движок, отсутствующая возможность или превышение заданного ограничения. Мягкие факторы - стадия, допустимая сложность, срок и соответствие приоритету - влияют на оценку, риск и порядок, но не должны маскироваться под физическую невозможность.")
add_paragraph(doc, "Пошагово расчёт выполняется следующим образом:")
add_bullets(doc, [
    "нормализуются входы профиля и удаляются дубликаты функций и платформ;",
    "из каталога выбираются опубликованные методы с указанным источником;",
    "применяются фильтры совместимости и формируются объяснения исключений;",
    "для допустимых методов формируется матрица критериев;",
    "матрица ранжируется по TOPSIS с весами, зависящими от приоритета пользователя;",
    "к результату добавляются риски стадии, связи с движком, альтернативы и устойчивость ранга;",
    "для текущей корзины отдельно рассчитываются выбранные методы, конфликты, зависимости и сводный профиль нагрузки.",
])
add_code_listing(doc, "Листинг 2.1 - Упрощённый алгоритм формирования результата", """
profile = normalize(profile_input)
basket = canonicalize(selected_methods)
input_key = hash(profile, basket)
if input_key != current_result_key:
    discard_result()
candidates = filter_catalog(profile)
ranked = topsis(candidates, profile.priority)
result = build_result(profile, ranked, basket, input_key)
return result
""")
add_paragraph(doc, "Листинг показывает логику взаимодействия компонентов, а не полный исходный код. В реальной реализации те же проверки разделены между backend-расчётом и состоянием frontend. После изменения профиля или корзины ответ старого запроса не должен возвращаться на экран.")

add_heading(doc, "2.4 Ранжирование и устойчивость решения", 2)
add_paragraph(doc, "TOPSIS применяется к матрице «альтернатива - критерий». Критерии делятся на выгодные и затратные, после чего для каждой допустимой альтернативы рассчитывается близость к идеальному решению. Формулы нормализации и коэффициента близости приведены в разделе 5, чтобы математический аппарат был собран в одном месте.")
add_paragraph(doc, "В текущем профиле учитываются производительность, сохранение качества, соответствие концепции, стоимость внедрения, штраф позднего внедрения, риск, соответствие ресурсам и уверенность. При одной альтернативе или одинаковых строках матрицы строгий порядок не имеет смысла. Разница менее 0,02 трактуется как группа практически равнозначных вариантов, а не как доказанное превосходство лидера.")
add_paragraph(doc, "Для проверки устойчивости веса каждого критерия по одному изменяются на ±10 %. Система показывает диапазон мест, которое может занимать метод. Такая процедура показывает чувствительность результата к политике выбора, но не превращает экспертные исходные оценки в измеренные данные.")

add_heading(doc, "2.5 Корзина и связи между решениями", 2)
add_paragraph(doc, "Корзина хранит один текущий набор решений и профиль проекта. Для набора проверяются конфликты, зависимости и взаимодополняющие связи. Конфликт означает, что решения нельзя одновременно применять в заявленной конфигурации. Зависимость означает наличие предварительного условия. Связь complement сама по себе не является измеренным численным бонусом: дополнительный эффект разрешено добавлять только при наличии обоснования совместного механизма.")
add_paragraph(doc, "Состав корзины нормализуется: дубликаты удаляются, порядок кодов не влияет на отпечаток входа. После изменения профиля или корзины старый результат сбрасывается, а ответ устаревшего запроса отбрасывается.")
add_paragraph(doc, "Сложность замены нельзя вывести разностью двух карточек: она зависит от базовой реализации. Может измениться конфигурация, формат данных и пайплайн контента или архитектурный слой. Поэтому универсальная «стоимость переписывания» не рассчитывается; для конкретной пары оставляется краткая карточка позднего внедрения и поле предметного анализа.")

add_heading(doc, "2.6 Модель нагрузки и аппаратной оценки", 2)
add_paragraph(doc, "Вычислительная модель разделяет активную сцену и объём мира. Активная сцена определяет работу на кадр, а масштаб мира дополнительно влияет на потоковую загрузку, резидентную память и общий объём контента. Формулы бюджета кадра и распределения частот приведены в разделе 5. Оценка оборудования остаётся ориентировочной и одновременно проверяет производительность, память и обязательные аппаратные возможности.")

add_heading(doc, "2.7 Выводы по разделу 2", 2)
add_paragraph(doc, "Модель связывает входы профиля, каталог решений, фильтрацию, ранжирование и корзину в единую последовательность. Основное методическое требование - не превращать экспертный коэффициент или совпадение двух записей в доказанную причинную связь. Там, где данных недостаточно, система должна показывать ограничение и направлять пользователя к последующей проверке в движке.")

# Section 3.
add_heading(doc, "3 АРХИТЕКТУРА И РЕАЛИЗАЦИЯ СИСТЕМЫ", 1, page_break=True)
add_heading(doc, "3.1 Общая архитектура", 2)
add_paragraph(doc, "Проект реализован как локальное веб-приложение. Пользовательский интерфейс на React и TypeScript отправляет запросы к FastAPI. Backend валидирует профиль, обращается к репозиториям каталога и передаёт данные вычислительным сервисам. Результат возвращается в едином контракте и отображается в интерфейсе рекомендаций, нагрузки, оборудования и корзины.")
add_table(doc, "Таблица 3.1 - Компоненты архитектуры", ["Компонент", "Технологии / файлы", "Ответственность"], [42, 64, 64], [
    ("Frontend", "React, TypeScript, Vite", "анкета, экраны, корзина, сохранение текущего состояния"),
    ("API", "FastAPI, Pydantic", "валидация запросов, маршруты каталога и расчёта"),
    ("Доменное ядро", "rules.py, recommender.py", "применимость, риски, ранжирование и объяснения"),
    ("Расчёт ресурсов", "hardware.py", "стоимость кадра, память, оборудование и ограничения"),
    ("Данные", "SQLAlchemy, SQLite, Alembic", "каталог, связи, источники, миграции и целостность"),
    ("Наполнение", "seed/*.py, seed/data", "первичное наполнение опубликованного каталога"),
    ("Контроль качества", "pytest, Vitest, CI", "модульные, интеграционные и сборочные проверки"),
])
add_pipeline_figure(doc)
add_paragraph(doc, "В собранном frontend запросы к адресу /api проксируются на backend по адресу 127.0.0.1:8000. В production-сборке FastAPI может раздавать собранные статические файлы frontend. База данных по умолчанию хранится локальным SQLite-файлом; отдельная установка сервера базы данных не требуется.")

add_heading(doc, "3.2 Предметные данные и публикация", 2)
add_paragraph(doc, "Основные сущности базы данных - игровая функция, метод, движок, инструмент движка, связь метода с инструментом, связь между методами, аппаратная запись CPU/GPU и журнал проверки. Для публичного расчёта используются опубликованные записи с указанным источником. Состояние публикации отделено от исследовательских черновиков.")
add_table(doc, "Таблица 3.2 - Текущее наполнение локального каталога", ["Объект", "Количество", "Статус / оговорка"], [52, 28, 90], [
    ("Игровые функции", "36", "текущий каталог функций"),
    ("Методы", "111", "все имеют статус published и источник в текущей базе; это не заменяет проверку содержания"),
    ("Движки / инструменты", "7 / 70", "связь с конкретной технологией показывается отдельно"),
    ("Связи метод - инструмент", "473", "тип связи может быть прямым, частичным, альтернативным или диагностическим"),
    ("Связи между методами", "20", "конфликты, зависимости и дополнения требуют предметного толкования"),
    ("CPU / GPU", "51 / 79", "каталог референсных аппаратных записей"),
    ("Ошибки целостности", "0", "текущее чтение validation_issues из локальной базы"),
])
add_paragraph(doc, "Из 111 методов 48 записей в текущей базе ссылаются на Wikipedia. Это не делает запись недействительной автоматически, но означает, что статус published следует понимать как статус наполнения каталога, а не как независимое доказательство каждой причинной связи. Четыре записи не имеют прямого численного эффекта по ресурсам; для них требуется либо объяснение области эффекта, либо уточнение данных.")

add_heading(doc, "3.3 API и единый результат расчёта", 2)
add_table(doc, "Таблица 3.3 - Основные API-операции", ["Маршрут", "Назначение", "Результат"], [45, 61, 64], [
    ("GET /api/health", "проверка готовности", "версия и состояние базы"),
    ("POST /api/recommend", "рекомендации", "методы, ранги, причины, риски и корзина"),
    ("POST /api/load-profile", "сводная нагрузка корзины", "CPU/GPU/RAM/VRAM и качественные эффекты"),
    ("POST /api/hardware-estimate", "оборудование", "референсные CPU/GPU, память и оговорки"),
    ("GET /api/catalog/methods", "каталог методов", "паспорта опубликованных решений"),
    ("GET /api/catalog/engines", "каталог движков", "версии и инструменты"),
    ("GET /api/docs", "интерактивная документация", "схема API Swagger"),
])
add_paragraph(doc, "Единый ответ рекомендаций включает отпечаток входа, версию алгоритма и версию каталога. Отпечаток строится по нормализованным данным профиля и корзины, поэтому результат можно проверить на соответствие текущему состоянию. Внутри ответа находятся также профиль нагрузки, аппаратная оценка и список применённых или исключённых методов.")

add_heading(doc, "3.4 Состояние frontend и сохранение текущего проекта", 2)
add_paragraph(doc, "Frontend хранит одну текущую анкету и одну текущую корзину в sessionStorage под ключом gamedev_dss_project_v1. История версий намеренно не включена в границы текущей версии: сравнение выполняется через изменение текущего набора и повторный расчёт. При изменении входов результат инвалидируется. Для сетевого запроса создаётся AbortController, а ответ проверяется по отпечатку запрошенного состояния.")
add_paragraph(doc, "Такое решение соответствует задаче поддержки одного текущего черновика. Если в дальнейшем понадобится история архитектурных вариантов, её следует проектировать отдельно, чтобы не смешать журнал изменений с источником текущего расчёта.")

add_heading(doc, "3.5 Запуск на Windows и Linux", 2)
add_paragraph(doc, "Текущий автоматический сценарий запуска оформлен для Windows через start.bat и отдельные run.bat. Он создаёт виртуальное окружение, устанавливает зависимости, запускает backend и frontend, проверяет /api/health и наличие React-контейнера, после чего открывает браузер. Полный сценарий быстрого старта требует финальной проверки после очистки проекта; место для результата оставлено в приложении В.")
add_paragraph(doc, "Код приложения рассчитан также на запуск на Linux. FastAPI, SQLAlchemy, SQLite, React, TypeScript и Vite не требуют Windows API. На Linux используются каталоги .venv/bin вместо .venv\\Scripts, а текущие batch-файлы заменяются shell-скриптом либо ручным запуском двух сервисов. CI использует ubuntu-latest и выполняет тесты, миграции, проверку типов и production-сборку, что подтверждает переносимость цепочки сборки. Отдельную проверку запуска на реальной Linux-системе следует добавить после подготовки start.sh.")

add_heading(doc, "3.6 Выводы по разделу 3", 2)
add_paragraph(doc, "Архитектура разделяет интерфейс, API, вычислительное ядро, каталог и расчёт ресурсов. Сохранение текущего состояния и отпечаток результата защищают от показа ответа для старого профиля. Основной проект завершён по функциональной структуре; последующие действия относятся к финальной упаковке, проверке запуска и сокращению тестового набора.")

# Section 4.
add_heading(doc, "4 ТЕХНИЧЕСКИЕ МЕТОДЫ И ИХ ВЛИЯНИЕ НА ИГРУ", 1, page_break=True)
add_heading(doc, "4.1 Механизмы, эффекты и ограничения методов", 2)
add_paragraph(doc, "Технический метод рассматривается через заменяемую работу, механизм, новые расходы и способ проверки. Нельзя считать оптимизацию безусловной скидкой: ускорение одной стадии может добавить вычислительный проход, память, время сборки, задержку или ограничения качества. Например, рост occupancy помогает скрывать задержки, но сам по себе не гарантирует ускорение и может ухудшить поведение кэша [2].")
add_paragraph(doc, "В каталоге выделены семейства culling; LOD/HLOD и impostors; batching и instancing; streaming, virtual texturing и compression; caching и pooling; baking и precompute; GPU offload и meshlets; upscaling, temporal и frame generation; tick/update budgeting. Для каждого семейства карточка должна фиксировать сокращаемую работу, новые затраты, условия и проверку. Формулировки описывают механизм и не задают универсальный процент ускорения.")

add_heading(doc, "4.1.1 Управление видимостью и детализацией", 3)
add_paragraph(doc, "Culling исключает из дальнейшей обработки объекты, которые не должны влиять на текущий вид. LOD выбирает представление с меньшей детализацией по расстоянию или экранной ошибке, а HLOD объединяет элементы иерархии. Импостор заменяет сложную геометрию подготовленным представлением. Все варианты требуют проверки видимости, расстояний перехода, теней, силуэта и стоимости подготовки структур. В небольшой сцене стоимость самого механизма может оказаться сопоставимой с выигрышем, поэтому решение должно подтверждаться измерением конкретной сцены.")

add_heading(doc, "4.1.2 Batching и инстансинг", 3)
add_paragraph(doc, "Batching объединяет объекты или команды с совместимыми материалами и состояниями, а instancing использует одну геометрию для многих экземпляров. Ограничения материалов, шейдеров и видимости определяют, действительно ли уменьшается работа. Документация Unity отдельно указывает, что dynamic batching переносит подготовку на CPU и не всегда выгоден [3]. Документация Godot также подчёркивает, что объединённая геометрия может потерять отдельное отсечение [4]. Поэтому в карточке нужно фиксировать не просто «меньше draw calls», а условия, при которых это достигается.")

add_heading(doc, "4.1.3 Потоковая загрузка, виртуальное текстурирование и сжатие", 3)
add_paragraph(doc, "Потоковая загрузка поддерживает в памяти только часть мира и подготавливает следующую часть по запросу. Виртуальное текстурирование работает со страницами или плитками, а кэш выбирает резидентные данные. Управление residency и загрузочными ресурсами является отдельной задачей: документация Direct3D 12 разделяет резидентность ресурсов [5] и временные ресурсы загрузки [6]. Поэтому в модели разделены резидентный объём, транзитный пул, декомпрессия и возможный пик памяти.")
add_paragraph(doc, "Блочное сжатие текстур уменьшает полезную нагрузку, но итоговый размер зависит от формата, выравнивания, размещения и временных копий. Для BC-форматов исходным основанием служит описание блоков 4x4 [7]. Сжатие нельзя без дополнительных данных превращать в фиксированную экономию RAM, VRAM и CPU одновременно.")

add_heading(doc, "4.1.4 Предвычисление, кэширование и расписание", 3)
add_paragraph(doc, "Baking переносит часть вычислений из времени исполнения в этап подготовки контента. Это может уменьшить runtime-работу, но добавляет данные и ограничивает изменения сцены. Кэширование и pooling уменьшают повторные аллокации или вычисления, однако удерживают объекты и требуют корректной инвалидации. Снижение частоты AI или физики должно рассматриваться отдельно от частоты вывода кадра: в движках фиксированные обновления могут накапливаться и выполняться несколькими шагами [8].")
add_paragraph(doc, "Если оптимизация переносит работу на GPU, требуется учитывать подготовку команд, память и синхронизацию. Нельзя делать вывод о выгоде только по тому, что операция выполняется параллельно. Для будущего измерительного протокола полезны захват времени кадра и инструментированное профилирование; в проекте PresentMon [9] и Tracy рассматриваются как внешние кандидаты, а не встроенные зависимости.")

add_heading(doc, "4.1.5 Трассировка лучей, материалы и постобработка", 3)
add_paragraph(doc, "Трассировка лучей состоит как минимум из подготовки структур ускорения и traversal. В руководстве Unreal Engine различаются обновления BLAS и построение TLAS, а затраты связываются с потоками рендера и GPU [10]. Поэтому RT не должен учитываться как одна универсальная скидка или фиксированное число миллисекунд для любой игры.")
add_paragraph(doc, "Апскейлинг снижает внутреннее разрешение части пиксельных стадий и добавляет собственный проход восстановления. Генерация кадров имеет отдельную стоимость синтеза и не должна уменьшать стоимость уже отрисованных кадров. Для материалов важны конкретные вычисления шейдера и приближения: документация Filament показывает, что даже отдельные параметры PBR имеют собственную обработку и ограничения [11].")
add_paragraph(doc, "Аппаратно-зависимое сжатие может выполняться не тем устройством, которое ожидает разработчик. Например, DirectStorage допускает CPU fallback для декомпрессии [12]. Поэтому в паспорте следует указывать ресурс исполнения и не переносить свойство Windows API на Linux без отдельной проверки.")

add_heading(doc, "4.2 Геймплейные решения как технические требования", 2)
add_table(doc, "Таблица 4.1 - Связь геймплейного требования и архитектуры", ["Геймплейное решение", "Технические последствия", "Ограничение вывода"], [40, 75, 55], [
    ("Разрушаемость окружения", "динамическая геометрия, физические тела, обновление коллизий, сохранение состояния, возможная репликация", "сам факт разрушаемости не задаёт число объектов и частоту событий"),
    ("Большой мир", "стриминг уровней и страниц, HLOD, память, размер контента, I/O", "масштаб мира не равен числу одновременно активных сущностей"),
    ("Большое количество NPC", "AI, поиск пути, анимация, perception, LOD, расписание обновлений", "количество NPC не определяет сложность поведения без модели задач"),
    ("Мультиплеер", "авторитет, репликация, сериализация, топология, bandwidth, детерминизм", "число игроков не равно сетевой нагрузке без частоты и объёма состояния"),
    ("Высокое разрешение и FPS", "бюджет кадра и пиксельные стадии рендера", "FPS не является гарантией, если нет трассы конкретной игры"),
])
add_paragraph(doc, "Разделение активной сцены и объёма мира особенно важно для открытого мира. При одинаковом количестве активных объектов размер мира может увеличить резидентную память и стоимость потоковой загрузки, но не обязан увеличивать стоимость каждого кадра. Это правило используется и в текущей реализации расчёта.")

add_heading(doc, "4.3 Интеграция, совместимость и стоимость замены", 2)
add_paragraph(doc, "Методы влияют друг на друга через общие данные и этапы пайплайна: LOD может конфликтовать с ручным объединением геометрии, virtual texturing связано с residency и streaming pool, baking ограничивает динамику, а frame generation требует отдельного анализа базовой и отображаемой частоты. Общая цель или общий ресурс ещё не доказывают конфликт.")
add_paragraph(doc, "В карточке разделяются первоначальное и позднее внедрение, совместимость и фактическая стоимость перехода. Последняя зависит от базовой архитектуры и без анализа конкретной пары не рассчитывается: изменение параметра означает настройку, формата данных - переработку контента, архитектурного слоя - возможную существенную переработку. При нехватке сведений используется «требует анализа» или пустое поле. Исключённые решения уже определены границами предметной области в разделе 1; универсальные утверждения об ускорении не используются.")

# Section 5.
add_heading(doc, "5 МАТЕМАТИЧЕСКИЕ МОДЕЛИ, ПРОВЕРКА И РЕЗУЛЬТАТЫ", 1, page_break=True)
add_heading(doc, "5.1 Модель ресурсной нагрузки", 2)
add_paragraph(doc, "Вычислительная модель разделяет активную сцену и объём мира. Активная сцена определяет работу на кадр, а масштаб мира дополнительно влияет на потоковую загрузку, резидентную память и общий объём контента. Такое разделение предотвращает повторное умножение стоимости кадра только из-за размера мира.")
add_equation(doc, "Bкадр = 1000 / fрендер", "(5.1)")
add_paragraph(doc, "Бюджет кадра определяется частотой реально отрисованных кадров. Физика и AI имеют собственные частоты обновления. Их вклад распределяется по кадрам по правилу:")
add_equation(doc, "Lкадр = Lтакт × fтакт / fрендер", "(5.2)")
add_paragraph(doc, "Для CPU отдельно учитываются последовательный критический путь и условно распараллеливаемая часть. Для GPU разделяются растеризация, пиксельные стадии, вычисления и трассировка лучей. RAM и VRAM представляются как сумма поименованных компонентов, а повторяющаяся экономия одного и того же компонента не должна учитываться дважды.")
add_paragraph(doc, "Оценка оборудования выбирает минимальный класс из каталога, одновременно проверяя вычислительные индексы, объём памяти и обязательные аппаратные возможности. Это расчётная оценка, а не измерение конкретной игры.")

add_heading(doc, "5.2 Ранжирование методом TOPSIS", 2)
add_paragraph(doc, "TOPSIS применяется к матрице «альтернатива - критерий». Критерии делятся на выгодные, где большее значение предпочтительно, и затратные, где предпочтительно меньшее значение. После нормализации строки умножаются на веса, зависящие от приоритета пользователя.")
add_equation(doc, "rᵢⱼ = xᵢⱼ / √(Σᵢ xᵢⱼ²)", "(5.3)")
add_paragraph(doc, "Для каждого критерия строятся положительное и отрицательное идеальные решения. Коэффициент близости альтернативы к идеальному решению вычисляется через евклидовы расстояния:")
add_equation(doc, "Cᵢ = dᵢ⁻ / (dᵢ⁺ + dᵢ⁻)", "(5.4)")
add_paragraph(doc, "Чем выше Cᵢ, тем ближе вариант к положительному идеалу. При одной альтернативе или одинаковых строках матрицы строгий порядок не имеет смысла. Разница менее 0,02 трактуется как группа практически равнозначных вариантов, а не как доказанное превосходство лидера. Для проверки устойчивости веса критериев по одному изменяются на ±10 %.")
add_paragraph(doc, "Расчётный пример с условными значениями не является измерением производительности: он нужен только для проверки направления формулы. Фактические значения текущей модели приведены в таблице чувствительности ниже.")

add_heading(doc, "5.3 Организация проверки", 2)
add_paragraph(doc, "Проверка системы состоит из трёх уровней. Модульные тесты проверяют отдельные правила и формулы. Интеграционные тесты отправляют реальные запросы к FastAPI и проверяют связность каталога, расчёта и базы. Frontend-тесты проверяют состояние интерфейса, сохранение, инвалидирование результата и обработку устаревших ответов. Сквозной smoke-check обращается к запущенному сервису и подтверждает работу основных пользовательских сценариев.")
add_paragraph(doc, "В текущем проекте накоплен расширенный набор проверок. Для учебного отчёта не требуется перечислять каждую проверку: основной текст фиксирует свойства, которые должны быть подтверждены приёмочным набором, а полный набор сохраняется как регрессионный материал.")

add_heading(doc, "5.4 Чувствительность к изменению параметров профиля", 2)
add_paragraph(doc, "В качестве базового сценария использован 3D-проект с открытым миром, высоким качеством, разрешением 1440p, целевой частотой 60 FPS и выбранными функциями потоковой загрузки и толпы NPC. Значения индексов являются внутренними относительными величинами текущей модели, а RAM/VRAM и draw calls - расчётными оценками.")
add_table(doc, "Таблица 5.1 - Изменение результата при изменении параметров профиля", ["Сценарий", "CPU index", "GPU index", "RAM, ГБ", "VRAM, ГБ", "Draw calls"], [43, 23, 23, 25, 25, 31], [
    ("База: 60 FPS, 1440p", "0.3057", "0.3216", "14.9", "10.9", "40535"),
    ("30 FPS", "0.1529", "0.1608", "14.9", "10.9", "40535"),
    ("144 FPS", "0.7337", "0.7719", "14.9", "10.9", "40535"),
    ("720p", "0.3057", "0.1608", "14.4", "7.7", "40535"),
    ("2160p", "0.3057", "0.5513", "15.7", "15.5", "40535"),
    ("Масштаб small", "0.3057", "0.3216", "12.3", "8.4", "27419"),
    ("Масштаб very_large", "0.3057", "0.3216", "16.4", "12.3", "47822"),
])
add_paragraph(doc, "Результат соответствует ожидаемому направлению модели. Повышение целевой частоты уменьшает бюджет одного кадра и повышает относительные индексы требований. Повышение разрешения влияет прежде всего на GPU и целевые буферы. Изменение масштаба мира при неизменной активной сцене увеличивает память и объём контента, но не изменяет CPU/GPU-стоимость кадра. Это различие является принципиальным: иначе размер мира ошибочно считался бы повторной работой на каждом кадре.")
add_paragraph(doc, "Для финальной версии следует повторить таблицу после изменения каталога или коэффициентов и заменить значения в отмеченном месте, если они изменятся.")

add_heading(doc, "5.5 Изменение и сохранение корзины", 2)
add_paragraph(doc, "Для базового профиля было проверено добавление метода «Временное масштабирование изображения». При пустой корзине выбранных методов нет. После добавления метод появляется в selected_methods, отпечаток входа изменяется, а сводный профиль нагрузки получает отдельные эффекты метода.")
add_table(doc, "Таблица 5.2 - Пример реакции на изменение корзины", ["Состояние", "Изменение результата", "Интерпретация"], [45, 65, 57], [
    ("Пустая корзина", "selected_methods пуст; input_key соответствует профилю", "расчёт только по профилю и общим правилам"),
    ("Добавлено temporal_upscaling", "GPU raw -0.0272; RAM +0.0067; VRAM +0.0275", "снижается часть пиксельной стоимости, но добавляются буферы и проход восстановления"),
    ("Изменён порядок кодов", "отпечаток не меняется", "набор, а не порядок выбора является входом"),
    ("Метод заменён другим", "selected_methods и input_key меняются", "результат строится для нового набора; состав рекомендаций может остаться тем же"),
])
add_paragraph(doc, "Изменение корзины не обязано менять первые места TOPSIS. В сценарии изменились selected_methods, отпечаток и сводная нагрузка, а топ-5 остался тем же: корзина описывает выбранные решения, а кандидаты ранжируются отдельно.")
add_paragraph(doc, "Корзина хранится в черновике браузера и после обновления восстанавливает профиль и набор. Изменение профиля или корзины сбрасывает предыдущий ответ, а поздний ответ старого запроса отбрасывается. Трудоёмкость замены зависит от конкретной пары методов и оставляется для предметного анализа.")

add_heading(doc, "5.6 Ограничения достоверности и аппаратной оценки", 2)
add_paragraph(doc, "Проверки подтверждают, что ограничения RAM и VRAM обрабатываются как жёсткие условия: при слишком малом пределе система сообщает о нарушении. Если расчёт превышает доступный каталог, вместо случайного выбора возвращается признак exceeds_catalog. Аппаратные записи проверяются одновременно по производительности, памяти и обязательным возможностям, включая поддержку трассировки лучей.")
add_paragraph(doc, "Оценка не содержит поля гарантированного FPS и сопровождается оговорками об экспертных коэффициентах, неизвестных входах и необходимости измерения на конкретной сцене. Это ограничение является частью корректного результата, а не отказом от расчёта.")

add_heading(doc, "5.7 Текущее состояние тестов и сборки", 2)
add_table(doc, "Таблица 5.3 - Результаты контрольного прогона текущей версии", ["Область", "Результат", "Статус для отчёта"], [48, 52, 67], [
    ("Backend pytest", "452 passed, 2 skipped", "текущая расширенная регрессия; обновить после финального сокращения"),
    ("Frontend Vitest", "34 passed", "текущая проверка store и компонентов; обновить после чистки"),
    ("TypeScript", "без ошибок", "проверка типов пройдена"),
    ("Production build", "сборка завершена", "проверка frontend пройдена"),
    ("Smoke-check", "85/85 по текущему сценарию", "повторить после финальной проверки быстрого старта"),
    ("CI", "Ubuntu: тесты, миграции, типы, сборка", "подтверждает воспроизводимость цепочки"),
])
add_paragraph(doc, "Для учебного проекта проверки разделяются на основной приёмочный набор и расширенную регрессию. Основной набор подтверждает запуск, изменение параметра профиля, RAM/VRAM-ограничения, корзину и сохранение, одну связь решений, отсутствие обещания FPS и сборку; остальные проверки остаются регрессионным материалом.")
add_placeholder(add_paragraph(doc, "", first_line=False), "[После очистки указать финальное количество приёмочных и расширенных тестов.]")

add_paragraph(doc, "В проекте нет независимой трассы CPU/GPU, измеренного FPS, пиков памяти и стенда для подтверждения универсальных коэффициентов. Поэтому числа называются расчётными индексами и ориентировочными оценками. Поля source_url и published подтверждают заполненность карточки, но не индивидуальную проверку каждой связи и влияния; стоимость перехода без описания базовой реализации остаётся пустой или условной.")
add_paragraph(doc, "Проект работает локально, не содержит встроенного runtime-профилировщика и не применяет решения автоматически. На Windows используется batch-сценарий; Linux требует отдельного shell-скрипта.")

add_paragraph(doc, "Итог раздела 5: расчёт реагирует на значимые параметры профиля, корзина меняет выбранный набор и сводный результат, а устаревшие ответы не должны подменять текущее состояние. После чистки тестов обновляются количественные показатели и ссылки на финальные сценарии.")

# Conclusion.
add_heading(doc, "ЗАКЛЮЧЕНИЕ", 1, page_break=True)
add_paragraph(doc, "В рамках проекта разработана локальная информационная система поддержки принятия решений по технической оптимизации разработки компьютерных игр. Система связывает профиль проекта, игровые функции, методы реализации, инструменты движков, ограничения и ориентировочную аппаратную оценку. Основная функциональность проекта реализована и объединена в единый пользовательский сценарий.")
add_paragraph(doc, "В отчёте сформулирована граница технического анализа. В неё входят решения, влияющие на архитектуру, вычисления и ресурсы: разрушаемость, масштаб мира, количество NPC, мультиплеер, рендеринг, память, потоковая загрузка, физика, AI и сетевые подсистемы. Монетизация, дата выхода, сюжетные развилки и количество концовок исключены из расчёта, поскольку не задают в данной модели проверяемую техническую нагрузку.")
add_paragraph(doc, "Реализованный расчёт использует фильтрацию применимости, TOPSIS, анализ устойчивости, единый профиль нагрузки и подбор референсного класса оборудования. Показано, что целевой FPS и разрешение меняют соответствующие индексы, а размер мира при одинаковой активной сцене влияет прежде всего на память и потоковую загрузку. Оценка оборудования остаётся ориентировочной и не является гарантией FPS.")
add_paragraph(doc, "Проверка корзины показала, что изменение состава набора меняет отпечаток входа, список выбранных методов и сводное влияние ресурсов. При этом рейтинг свободных кандидатов может не измениться, что не является ошибкой. Сложность перехода от уже реализованного решения к другому зависит от базовой архитектуры, поэтому она не заполняется универсальным числом без дополнительного анализа.")
add_paragraph(doc, "Основные оставшиеся действия относятся к финальной редакции: проверить формулировки карточек и источников, удалить неподтверждённые причинные выводы, завершить анализ взаимодействий методов, сократить основной набор тестов, проверить быстрый старт на Windows и Linux и обновить количественные показатели в отмеченных местах. Эти действия не меняют того, что основная разработка системы завершена.")
# References.
add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 1, page_break=True)
add_source(doc, 1, "GAMORRA: A GPU and CPU Performance Modeling Framework for Games (arXiv 2204.11025, 2022)", "https://arxiv.org/pdf/2204.11025", "использовано для разделения параметров, модели и измерительного протокола")
add_source(doc, 2, "Guthmann F. Occupancy explained. AMD GPUOpen, 2023; обновление 2024", "https://gpuopen.com/learn/occupancy-explained/", "использовано для ограничения вывода об occupancy")
add_source(doc, 3, "Draw call batching. Unity Manual 6000.0", "https://docs.unity3d.com/6000.0/Documentation/Manual/DrawCallBatching.html", "использовано для условий batching")
add_source(doc, 4, "GPU optimization. Godot Engine documentation 4.4", "https://docs.godotengine.org/en/4.4/tutorials/performance/gpu_optimization.html", "использовано для batching, материалов и pixel cost")
add_source(doc, 5, "D3D12 Residency. Microsoft Learn", "https://learn.microsoft.com/en-us/windows/win32/direct3d12/residency", "использовано для различения резидентности и рабочего набора")
add_source(doc, 6, "Uploading resources. Microsoft Learn", "https://learn.microsoft.com/en-us/windows/win32/direct3d12/uploading-resources", "использовано для описания временных ресурсов загрузки")
add_source(doc, 7, "Texture block compression in Direct3D 11. Microsoft Learn", "https://learn.microsoft.com/en-us/windows/win32/direct3d11/texture-block-compression-in-direct3d-11", "использовано для ограничения расчёта размера текстур")
add_source(doc, 8, "Fixed updates. Unity Manual 6000.0", "https://docs.unity3d.com/6000.0/Documentation/Manual/fixed-updates.html", "использовано для разделения частоты такта и частоты кадра")
add_source(doc, 9, "PresentMon. Intel GameTechDev", "https://github.com/GameTechDev/PresentMon", "рассматривается как внешний инструмент измерения времени кадра")
add_source(doc, 10, "Ray Tracing Performance Guide in Unreal Engine 5.5. Epic Games", "https://dev.epicgames.com/documentation/en-us/unreal-engine/ray-tracing-performance-guide-in-unreal-engine?application_version=5.5", "использовано для разделения BLAS/TLAS и RT-проходов")
add_source(doc, 11, "Filament: Physically Based Rendering. Google", "https://google.github.io/filament/Filament.md.html", "использовано для конкретизации обработки материалов и приближений")
add_source(doc, 12, "DirectStorage compression support. Microsoft Learn", "https://learn.microsoft.com/en-us/windows/win32/dstorage/dstorage/ne-dstorage-dstorage_compression_support", "использовано для ограничения вывода о декомпрессии")

# Appendices.
add_heading(doc, "ПРИЛОЖЕНИЕ А. СВОДКА ТЕКУЩЕГО СОСТОЯНИЯ ПРОЕКТА", 1, page_break=True)
add_paragraph(doc, "Приложение предназначено для обновления после финальной очистки. В основной текст отчёта следует переносить только те числа, которые подтверждены текущим прогоном и совпадают с каталогом, используемым приложением.")
add_table(doc, "Таблица А.1 - Состояние проекта на момент подготовки рабочей версии", ["Область", "Текущее состояние", "Место обновления"], [46, 74, 54], [
    ("Backend", "FastAPI + SQLAlchemy + SQLite; расчёт, каталог, миграции", "[версия / commit]"),
    ("Frontend", "React + TypeScript + Vite; 11 экранов по текущему README", "[уточнить итоговое число экранов]"),
    ("Методы", "111 published, все с source_url", "[проверить спорные карточки]"),
    ("Игровые функции", "36", "[обновить при изменении seed]"),
    ("Движки / инструменты", "7 / 70", "[обновить при изменении seed]"),
    ("Аппаратный каталог", "51 CPU / 79 GPU", "[обновить при изменении seed]"),
    ("Тесты", "452 backend; 34 frontend до финальной чистки", "[вписать основной и расширенный набор]"),
    ("Быстрый старт", "Windows batch; Linux ручной запуск / будущий shell-скрипт", "[вписать результат финального прогона]"),
])

add_heading(doc, "ПРИЛОЖЕНИЕ Б. КОМПАКТНЫЙ ПРИЁМОЧНЫЙ ПРОТОКОЛ", 1, page_break=True)
add_paragraph(doc, "Протокол предназначен для основного учебного прогона после сокращения расширенной регрессии. Каждый сценарий должен использовать чистую временную базу и фиксированный профиль.")
add_table(doc, "Таблица Б.1 - Минимальный набор приёмочных проверок", ["№", "Проверка", "Ожидаемый результат", "Факт"], [15, 58, 76, 21], [
    ("1", "Запуск и health-check", "backend отвечает, frontend открывается, база доступна", "[ ]"),
    ("2", "Один параметр профиля", "при изменении FPS или разрешения меняются соответствующие индексы", "[ ]"),
    ("3", "Разделение мира и активной сцены", "масштаб меняет память/контент, но не кадр при той же активной сцене", "[ ]"),
    ("4", "Аппаратные ограничения", "нарушение RAM/VRAM и отсутствие подходящего класса сообщаются явно", "[ ]"),
    ("5", "Корзина", "добавление/замена меняет selected_methods, input_key и сводный профиль", "[ ]"),
    ("6", "Сохранение", "после обновления восстанавливаются профиль и корзина; старый ответ сбрасывается", "[ ]"),
    ("7", "Совместимость", "конфликт или зависимость показывается с объяснением", "[ ]"),
    ("8", "Сборка", "typecheck и production build завершаются без ошибок", "[ ]"),
])
add_paragraph(doc, "Расширенные тесты, проверки безопасности, калибровочные прогоны и исторические контрпримеры следует оставить в проекте только при наличии практической пользы. В основной отчёт достаточно включить их классификацию и итоговый статус.")

add_heading(doc, "ПРИЛОЖЕНИЕ В. ПОЛЯ ДЛЯ ФИНАЛЬНОГО ОБНОВЛЕНИЯ", 1, page_break=True)
add_table(doc, "Таблица В.1 - Изменения, которые можно внести без перестройки отчёта", ["Поле", "Что добавить", "Готово"], [50, 99, 21], [
    ("Тесты", "итоговое число приёмочных и расширенных тестов после чистки", "[ ]"),
    ("Быстрый старт Windows", "дата, команда и результат запуска", "[ ]"),
    ("Быстрый старт Linux", "результат запуска через start.sh или ручные команды", "[ ]"),
    ("Каталог методов", "проверенные карточки и ссылки первичных источников", "[ ]"),
    ("Взаимодействия", "условия для конфликтов, зависимостей и complement", "[ ]"),
    ("Корзина", "конкретный пример замены решения и трудоёмкости перехода", "[ ]"),
    ("Профиль", "финальные значения таблицы чувствительности", "[ ]"),
    ("Титульный лист", "организация, кафедра, ФИО, группа, руководитель, город", "[ ]"),
    ("Использование ИИ", "указать факт и роль инструмента по требованиям преподавателя; итоговый текст проверить и изложить автором", "[ ]"),
])
add_placeholder(add_paragraph(doc, "", first_line=False), "[Поля в квадратных скобках предназначены для ручного обновления автором после финального прогона.]")

doc.save(OUT_DOCX)
for asset in _TEMP_ASSETS:
    try:
        asset.unlink(missing_ok=True)
    except OSError:
        pass
print(OUT_DOCX)
